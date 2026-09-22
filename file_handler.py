import os,base64,mimetypes
from pathlib import Path
import config
TEXT_EXTS={'.txt','.md','.json','.py','.js','.html','.css','.log','.xml','.yaml','.yml','.sql'}
IMAGE_EXTS={'.png','.jpg','.jpeg','.gif','.webp','.bmp'}
ALLOWED_EXTS={'.pdf','.docx','.csv','.xlsx','.xls'}|TEXT_EXTS|IMAGE_EXTS
def _truncate(text):
    return text if len(text)<=config.MAX_FILE_CHARS else text[:config.MAX_FILE_CHARS]+'\n\n[contenido recortado]'
def _check(path):
    p=Path(path)
    if not p.exists() or not p.is_file(): raise ValueError('Archivo no encontrado.')
    if p.stat().st_size>config.MAX_UPLOAD_MB*1024*1024: raise ValueError(f'Archivo supera {config.MAX_UPLOAD_MB} MB.')
    if p.suffix.lower() not in ALLOWED_EXTS: raise ValueError(f'Tipo no permitido: {p.suffix}')
def read_pdf(path):
    from pypdf import PdfReader
    return _truncate('\n'.join(f'--- Página {i} ---\n{p.extract_text() or ""}' for i,p in enumerate(PdfReader(path).pages,1)))
def read_docx(path):
    import docx
    d=docx.Document(path); out=[p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables: out.extend(' | '.join(c.text.strip() for c in r.cells) for r in t.rows)
    return _truncate('\n'.join(out))
def read_tabular(path):
    import pandas as pd
    df=pd.read_csv(path) if path.lower().endswith('.csv') else pd.read_excel(path)
    return _truncate(f'Filas: {len(df)}, Columnas: {len(df.columns)}\nColumnas: '+', '.join(map(str,df.columns))+'\n\n--- Primeras filas ---\n'+df.head(30).to_string(index=False))
def encode_image(path):
    mime=mimetypes.guess_type(path)[0] or 'image/png'
    return f"data:{mime};base64,"+base64.b64encode(Path(path).read_bytes()).decode()
def analyze_file(path):
    _check(path); ext=Path(path).suffix.lower(); name=Path(path).name
    if ext in IMAGE_EXTS: return 'image',encode_image(path)
    try:
        if ext=='.pdf': content=read_pdf(path)
        elif ext=='.docx': content=read_docx(path)
        elif ext in {'.csv','.xlsx','.xls'}: content=read_tabular(path)
        else: content=_truncate(Path(path).read_text(encoding='utf-8',errors='replace'))
    except Exception as e: content=f"[No se pudo leer '{name}': {e}]"
    return 'text',f"Contenido del archivo '{name}':\n{content}"
